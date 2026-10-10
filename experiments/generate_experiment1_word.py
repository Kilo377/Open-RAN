from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(
    r"E:\oran_slot_sim_outputs\paper_eval\experiment1_cross_scenario_transfer_targeted_7methods_v5_v15_combined"
)
FIGURE_DIR = ROOT / "figures_transfer"
OUT_PATH = ROOT / "Experiment1_Cross_Scenario_Orchestration_Detailed_Explanation.docx"
CSV_PATH = ROOT / "transfer_headline_results.csv"

FIGURES = [
    ("Figure 1", "Safety-first transfer score", FIGURE_DIR / "exp1_transfer_safety_first_score.png"),
    ("Figure 2", "Tail-robustness score", FIGURE_DIR / "exp1_transfer_tail_robustness_score.png"),
    ("Figure 3", "Safety violation rate", FIGURE_DIR / "exp1_transfer_safety.png"),
    ("Figure 4", "Near-RT safe-baseline fallback rate", FIGURE_DIR / "exp1_transfer_near_rt_fallback.png"),
    ("Figure 5", "Normalized intent loss", FIGURE_DIR / "exp1_transfer_intent_loss.png"),
    ("Figure 6", "Throughput trade-off", FIGURE_DIR / "exp1_transfer_throughput.png"),
]

METHOD_LABELS = {
    "fixed_plan": "Fixed Plan",
    "no_memory_planner": "No-Memory Planner",
    "success_only_memory": "Success-Only Memory",
    "orchestran_like_resource_aware": "OrchestRAN-like",
    "intent_hrl_like": "Intent-HRL-like",
    "oreo_like_service_graph": "OREO-like",
    "full_method": "Proposed Full Method",
}
METHOD_ORDER = [
    "fixed_plan",
    "no_memory_planner",
    "success_only_memory",
    "orchestran_like_resource_aware",
    "intent_hrl_like",
    "oreo_like_service_graph",
    "full_method",
]


def metric_text(row: dict[str, str], key: str, ndigits: int = 4) -> str:
    value = row.get(key, "")
    if value in (None, ""):
        return "N/A"
    try:
        return f"{float(value):.{ndigits}f}"
    except ValueError:
        return str(value)


def set_run_font(run, *, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    set_run_font(run, size=9, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_formula(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Formula"]
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size in [("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        styles[style_name].font.name = "Microsoft YaHei"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[style_name].font.size = Pt(size)
    try:
        formula_style = styles.add_style("Formula", 1)
    except ValueError:
        formula_style = styles["Formula"]
    formula_style.font.name = "Consolas"
    formula_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    formula_style.font.size = Pt(9)


def add_title(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Experiment 1 Detailed Explanation\nCross-Scenario Multi-xApp Orchestration")
    set_run_font(run, size=18, bold=True)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Comparison with Existing Orchestration and Intent-Driven Control Methods")
    set_run_font(run, size=11)
    run.italic = True


def add_result_table(doc: Document, rows_by_method: dict[str, dict[str, str]]) -> None:
    doc.add_heading("4. 数值结果总表", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Method", "Intent Loss ↓", "Safety Rate ↓", "Throughput ↑", "BLER ↓", "Safety-first ↓", "Tail Robust ↓"]
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        set_cell_shading(table.rows[0].cells[index], "D9EAF7")

    for method in METHOD_ORDER:
        row = rows_by_method[method]
        cells = table.add_row().cells
        values = [
            METHOD_LABELS[method],
            metric_text(row, "transfer_mean_intent_loss_normalized"),
            metric_text(row, "transfer_mean_safety_violation_rate"),
            metric_text(row, "transfer_mean_throughput_mbps"),
            metric_text(row, "transfer_mean_mean_bler"),
            metric_text(row, "transfer_safety_first_score"),
            metric_text(row, "transfer_tail_robustness_score"),
        ]
        for index, value in enumerate(values):
            set_cell_text(cells[index], value, bold=(method == "full_method"))
            if method == "full_method":
                set_cell_shading(cells[index], "E2F0D9")

    doc.add_paragraph(
        "表中可以看到，Proposed Full Method 的 safety violation rate、safety-first score 和 "
        "tail-robustness score 均为最优。它的 throughput 不是最高，这是安全优先编排带来的 trade-off："
        "吞吐型 baseline 通过更激进动作获得更高吞吐，但安全违反率明显更高。"
    )


def add_figure_explanations(doc: Document) -> None:
    doc.add_heading("5. 图像说明", level=1)
    for figure_id, title, figure_path in FIGURES:
        doc.add_heading(f"{figure_id}. {title}", level=2)
        if figure_path.exists():
            doc.add_picture(str(figure_path), width=Inches(6.2))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            doc.add_paragraph(f"[Missing figure: {figure_path}]")

        if "safety_first" in figure_path.name:
            doc.add_paragraph(
                "该图是实验一最核心的总览图。Safety-first score 按公式 "
                "S_safety-first = L_intent + 0.35 SafetyRate + 0.05 BLER + 0.02 SwitchingCost - 0.001 Throughput "
                "计算，越低越好。Proposed Full Method 得分为 0.0298，显著低于 Intent-HRL-like 的 0.0530、"
                "No-Memory Planner 的 0.0551 和 OrchestRAN-like 的 0.0573。"
                "这说明 proposed method 在跨场景 multi-xApp orchestration 中获得了最好的安全-性能综合权衡。"
            )
        elif "tail_robustness" in figure_path.name:
            doc.add_paragraph(
                "该图评价 worst-case transfer behavior。Tail-robustness score 在 safety-first score 基础上加入 "
                "worst-case score 和跨场景标准差：S_tail = S_safety-first + 0.25 WorstCase + 0.10 StdCase。"
                "Proposed Full Method 为 0.0441，优于 Intent-HRL-like 的 0.0728，说明 proposed 不只是平均值好，"
                "而且在困难 transfer 场景中更稳。"
            )
        elif figure_path.name == "exp1_transfer_safety.png":
            doc.add_paragraph(
                "该图展示 safety violation rate。Proposed Full Method 为 0.0465，是七个方法中最低；"
                "Intent-HRL-like 为 0.1163，No-Memory Planner 为 0.1279。该结果说明 failed-case reflection、"
                "typed constraints 和 pressure-adaptive safety gate 能显著减少 unsafe / unhealthy slots。"
            )
        elif "near_rt_fallback" in figure_path.name:
            doc.add_paragraph(
                "该图展示 Near-RT safe-baseline fallback rate。Proposed Full Method 为 0.0521。较低 fallback 表明 "
                "Near-RT bounded tuning 大多数时间可以在 Non-RT 下发的安全区域内执行，而不是频繁退回保守基线。"
                "这支持论文中“LLM 不进入 Near-RT path，Near-RT 只做 deterministic bounded update”的设计。"
            )
        elif "intent_loss" in figure_path.name:
            doc.add_paragraph(
                "该图展示 normalized intent loss。Proposed Full Method 为 0.01587，与 OrchestRAN-like 持平，"
                "略高于 No-Memory Planner 和 Intent-HRL-like，但优于 Fixed Plan、OREO-like 和 Success-Only Memory。"
                "这说明 proposed 的主要优势不是单点降低 intent loss，而是在保持 intent loss 可接受的同时显著降低安全风险和 tail risk。"
            )
        elif "throughput" in figure_path.name:
            doc.add_paragraph(
                "该图展示 throughput trade-off。Fixed Plan、OREO-like 和 Success-Only Memory 的吞吐较高，但它们的 "
                "safety violation rate 也更高。Proposed Full Method 的吞吐不是最高，说明其目标是安全优先的跨场景编排，"
                "而不是无约束 throughput maximization。论文中建议将该图作为 trade-off / appendix 图，而不是主结论图。"
            )


def build_document() -> None:
    rows = list(csv.DictReader(CSV_PATH.open(encoding="utf-8")))
    rows_by_method = {row["variant_name"]: row for row in rows}

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    configure_styles(doc)
    add_title(doc)

    doc.add_paragraph(
        "本文档说明实验一的实验目的、对比方法、跨场景评价指标、各图含义和主要结论。"
        "当前结果来自 targeted cross-scenario validation：先在 dense hotspot + high mobility 场景建立执行记忆，"
        "再迁移到 interference-edge 与 hotspot-shift/blockage 两个 transfer 场景。"
    )

    doc.add_heading("1. 实验目标与核心问题", level=1)
    doc.add_paragraph(
        "实验一被重新定位为 Cross-Scenario Multi-xApp Orchestration。它不是只验证平均 KPI 是否提升，"
        "而是验证 proposed reflective orchestration 在未见场景迁移时是否比静态策略、无记忆 planner、"
        "成功记忆、资源感知编排和 intent-driven HRL 更稳定、更安全。"
    )
    for item in [
        "Q1: reflective memory 是否比固定策略或 no-memory planner 更适合跨场景 multi-xApp orchestration?",
        "Q2: proposed method 是否优于 OrchestRAN-like、Intent-HRL-like、OREO-like 等文献启发 baseline?",
        "Q3: failed-case reflection 与 pressure-adaptive Near-RT safety gate 是否能降低困难 transfer 场景中的安全风险和 tail failure?",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("2. 实验设置", level=1)
    doc.add_paragraph("场景设置：")
    for item in [
        "Warm-up / memory-building scenario: oran_medium_dense_hotspot_high_mobility.",
        "Transfer scenario 1: oran_medium_interference_edge，用于测试小区边缘用户与干扰突发下的泛化能力。",
        "Transfer scenario 2: oran_medium_hotspot_shift_blockage，用于测试热点迁移、阻塞和移动性压力下的泛化能力。",
        "Intent: mobility-hotspot-stability，即降低 ping-pong handover / handover failure，同时维持 cell-edge throughput 与可靠性。",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("对比方法：")
    method_descriptions = [
        ("Fixed Plan", "所有场景使用固定 xApp 组合，作为静态策略下界。"),
        ("No-Memory Planner", "只根据当前 intent、scenario 和 registry 选 xApp，不检索 execution memory。"),
        ("Success-Only Memory", "只保存成功案例，不保存失败案例，用于隔离 failed-case reflection 的贡献。"),
        ("OrchestRAN-like", "资源/时延/切换成本感知编排，但不使用 execution-outcome reflective memory。"),
        ("Intent-HRL-like", "两层 intent-driven HRL：高层选 xApp set，低层选参数或 arbitration action。"),
        ("OREO-like", "service-graph / xApp sharing 风格 baseline，偏部署覆盖和资源共享优化。"),
        (
            "Proposed Full Method",
            "success + failed memory、typed conflict taxonomy、deterministic repair、typed guidance、pressure-adaptive Near-RT bounded tuning。",
        ),
    ]
    for name, description in method_descriptions:
        doc.add_paragraph(f"{name}: {description}", style="List Bullet")

    doc.add_heading("3. 指标定义与公式", level=1)
    doc.add_paragraph(
        "实验一不只使用单一综合指标，而是分开报告 intent fulfillment、safety、Near-RT fallback、tail robustness "
        "和 throughput trade-off。其中 safety-first score 和 tail-robustness score 用于总结跨场景安全泛化能力，数值越低越好。"
    )

    doc.add_heading("3.1 Normalized Intent Loss", level=2)
    doc.add_paragraph(
        "对每个 KPI 约束 l，先计算 raw gap，再按 reference / tolerance / normalization value 归一化，clip 后平方惩罚。"
    )
    add_formula(doc, "denom_l = max(normalization_value_l, tolerance_l, |reference_l|, 1.0)")
    add_formula(doc, "gap_l = max(0, reference_l - actual_l) / denom_l          for max-type KPI")
    add_formula(doc, "gap_l = max(0, actual_l - reference_l) / denom_l          for min-type KPI")
    add_formula(doc, "gap_l = (actual_l - reference_l) / denom_l                for target-type KPI")
    add_formula(doc, "penalty_l = clip(gap_l, -2.0, 2.0)^2")
    add_formula(doc, "L_intent = sum_l w_l * penalty_l")
    doc.add_paragraph(
        "transfer_mean_intent_loss_normalized 是所有 transfer case 上 L_intent 的平均值。该指标越低，说明 KPI constraint violation 越小。"
    )

    doc.add_heading("3.2 Safety Violation Rate", level=2)
    add_formula(doc, "SafetyViolationRate = (# slots violating safety / health-gate constraints) / (# evaluated slots)")
    doc.add_paragraph(
        "该指标统计 slot 级别是否违反安全/健康门控条件，例如 BLER 超过 gate、队列压力过高、handover pressure 过高或 protected QoS 不满足。"
    )

    doc.add_heading("3.3 Safety-First Transfer Score", level=2)
    doc.add_paragraph("当前代码中的 transfer_safety_first_score 定义如下：")
    add_formula(doc, "S_safety-first = L_intent + 0.35 * SafetyRate + 0.05 * BLER + 0.02 * SwitchingCost - 0.001 * Throughput")
    doc.add_paragraph(
        "其中 Throughput 的负号表示吞吐越高越好，但权重较小，因为该实验一的核心主张是安全跨场景编排，"
        "不是单纯吞吐最大化。这个 score 越低越好。"
    )

    doc.add_heading("3.4 Tail-Robustness Score", level=2)
    doc.add_paragraph("为了避免只优化平均值，实验一还定义 worst-case 与跨场景波动惩罚：")
    add_formula(doc, "S_case = L_intent_case + 0.35 * SafetyRate_case + 0.05 * BLER_case + 0.02 * SwitchingCost_case - 0.001 * Throughput_case")
    add_formula(doc, "WorstCase = max_case S_case")
    add_formula(doc, "StdCase = std_case(S_case)")
    add_formula(doc, "S_tail = S_safety-first + 0.25 * WorstCase + 0.10 * StdCase")
    doc.add_paragraph("S_tail 越低，说明方法不仅平均表现好，而且在困难 transfer 场景中的 worst-case 风险更低。")

    doc.add_heading("3.5 Near-RT Safe-Baseline Fallback Rate", level=2)
    add_formula(doc, "FallbackRate = (# slots using safe_baseline_fallback) / (# evaluated slots)")
    doc.add_paragraph(
        "该指标越低，说明 Near-RT bounded tuning 更少退回保守安全基线，能够在安全区域内持续执行在线微调。"
    )

    add_result_table(doc, rows_by_method)
    add_figure_explanations(doc)

    doc.add_heading("6. 可直接放入论文的结果描述", level=1)
    doc.add_paragraph(
        "Figure X reports the safety-first transfer score across seven orchestration methods. "
        "The proposed reflective orchestration achieves the lowest score (0.0298), outperforming the strongest learning-based baseline, "
        "Intent-HRL-like orchestration (0.0530), as well as the no-memory planner (0.0551) and OrchestRAN-like resource-aware orchestration (0.0573). "
        "This improvement is mainly driven by a lower safety violation rate and better tail robustness under shifted O-RAN scenarios."
    )
    doc.add_paragraph(
        "Figure Y further evaluates tail robustness by combining the average safety-first score, the worst-case transfer score, "
        "and the cross-scenario standard deviation. The proposed method achieves the best tail-robustness score (0.0441), "
        "compared with 0.0728 for Intent-HRL-like orchestration. This indicates that failed-case reflection and pressure-adaptive "
        "Near-RT safety gating improve robustness beyond average KPI optimization."
    )
    doc.add_paragraph(
        "Although throughput-oriented baselines obtain higher average throughput, they also incur substantially higher safety violation rates. "
        "Therefore, the proposed method provides a better safety-performance trade-off for cross-scenario multi-xApp orchestration."
    )

    doc.add_heading("7. 写作注意事项", level=1)
    for item in [
        "正文主图建议使用 Figure 1, Figure 2, Figure 3, Figure 4；Figure 5 和 Figure 6 可作为辅助图或 appendix。",
        "不要声称 proposed 在所有单项 KPI 上都最好。更准确的表述是：proposed 在 safety-first、tail robustness 和 safety violation 上最优，同时保持 intent loss 与强 baseline 同一量级。",
        "吞吐图应解释为 trade-off：高吞吐 baseline 的安全违反率更高，因此不适合安全关键 O-RAN orchestration。",
        "正式大规模实验建议沿用这些指标，并扩展到 16 个 intent、更多 seeds 和更多 transfer scenario families。",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("8. 数据与图像来源", level=1)
    for item in [
        f"Result directory: {ROOT}",
        f"Headline metrics: {CSV_PATH}",
        f"Figures: {FIGURE_DIR}",
        "Metric implementation: experiments/run_experiment1_transfer.py and experiments/eval_metrics.py",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("9. Baseline 文献定位", level=1)
    references = [
        "OrchestRAN-like baseline: resource/timing-aware Open RAN intelligence orchestration, https://arxiv.org/abs/2201.05632",
        "Intent-HRL-like baseline: hierarchical reinforcement learning for intent-driven O-RAN control, https://arxiv.org/abs/2307.02754",
        "OREO-like baseline: xApp-based service graph orchestration and sharing, https://arxiv.org/abs/2405.18198",
        "Reflection mechanism inspiration: Reflexion language agents with verbal reinforcement learning, https://arxiv.org/abs/2303.11366",
        "Safe bounded control inspiration: Constrained Policy Optimization and safe RL shielding, https://arxiv.org/abs/1705.10528 and https://arxiv.org/abs/1708.08611",
    ]
    for reference in references:
        doc.add_paragraph(reference, style="List Bullet")

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Generated by Codex from local experiment outputs")
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
